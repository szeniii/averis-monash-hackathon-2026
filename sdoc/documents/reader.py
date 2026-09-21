"""Turn an attachment into text, whatever format it is.

One entry point -- read(inbox, path) -- that dispatches on file extension and
always returns a DocumentExtract. It never raises: a file it cannot read comes
back with read_ok=False, which the decision gate turns into NEEDS_REVIEW /
unreadable rather than a crash or a false mismatch.

Formats present in the dataset: 192 .txt, 28 .pdf, 22 .xlsx, 8 .docx.
The SI and BL of one email can be different formats.
"""
from __future__ import annotations

import re

from sdoc.schemas import DocumentExtract

# Documents that are NOT an SI or a BL. If we see one of these titles, the
# case is wrong_doc_type -- a human should look, and we must not "compare" it.
WRONG_DOC_MARKERS = {
    "COMMERCIAL INVOICE": "COMMERCIAL_INVOICE",
    "PACKING LIST": "PACKING_LIST",
    "CERTIFICATE OF ORIGIN": "CERTIFICATE_OF_ORIGIN",
    "DEBIT NOTE": "DEBIT_NOTE",
}

MIN_USEFUL_CHARS = 40  # below this we assume the read failed


def _role_from_path(path: str) -> str:
    name = path.upper()
    if re.search(r"_SI\.", name):
        return "SI"
    if re.search(r"_BL\.", name):
        return "BL"
    return "UNKNOWN"


def _detect_doc_type(text: str):
    head = text[:400].upper()
    for marker, label in WRONG_DOC_MARKERS.items():
        if marker in head:
            return label
    if "BILL OF LADING" in head:
        return "BILL_OF_LADING"
    if "SHIPPING INSTRUCTION" in head:
        return "SHIPPING_INSTRUCTION"
    return None


def _read_pdf(data: bytes) -> str:
    try:
        import pdfplumber
        import io
        out = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                out.append(page.extract_text() or "")
        return "\n".join(out)
    except Exception:
        return ""


def _read_docx(data: bytes) -> str:
    try:
        import docx
        import io
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:           # BLs put fields in table cells
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    except Exception:
        return ""


def _read_xlsx(data: bytes) -> str:
    try:
        import openpyxl
        import io
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    # "Label | value" so the field extractor sees the same
                    # shape it sees in the .txt documents.
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    except Exception:
        return ""


def read(inbox, path: str) -> DocumentExtract:
    """Read one attachment. Never raises."""
    doc = DocumentExtract(path=path, doc_role=_role_from_path(path))
    lower = path.lower()

    try:
        data = inbox.read_bytes(path)
    except Exception as exc:
        doc.read_ok = False
        doc.read_error = f"could not open file: {exc}"
        return doc

    if not data:
        doc.read_ok = False
        doc.read_error = "file is empty (0 bytes)"
        return doc

    if lower.endswith(".txt"):
        doc.read_method = "plain_text"
        doc.text = data.decode("utf-8", errors="replace")
    elif lower.endswith(".pdf"):
        doc.read_method = "pdf"
        doc.text = _read_pdf(data)
    elif lower.endswith(".docx"):
        doc.read_method = "docx"
        doc.text = _read_docx(data)
    elif lower.endswith(".xlsx"):
        doc.read_method = "xlsx"
        doc.text = _read_xlsx(data)
    else:
        doc.read_ok = False
        doc.read_error = f"unsupported file type: {path}"
        return doc

    if len(doc.text.strip()) < MIN_USEFUL_CHARS:
        # Scanned image-only PDF, or a corrupt/truncated file. This is the
        # hand-off point for OCR / a vision model later.
        doc.read_ok = False
        doc.read_error = (
            f"no usable text from {doc.read_method} "
            f"({len(doc.text.strip())} chars) - likely a scan or corrupt file"
        )
        return doc

    doc.doc_type_detected = _detect_doc_type(doc.text)
    return doc
